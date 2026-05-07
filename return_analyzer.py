# -*- coding: utf-8 -*-
"""
退货分析核心脚本 - 服务器版本
支持手动触发分析，输出JSON格式结果
"""
import pandas as pd
import numpy as np
import json
import sys
import os
from datetime import datetime

def clean_for_json(obj):
    """清理数据中的NaN值，使其JSON兼容"""
    if isinstance(obj, dict):
        return {k: clean_for_json(v) for k, v in obj.items()}
    elif isinstance(obj, list):
        return [clean_for_json(v) for v in obj]
    elif isinstance(obj, float) and (np.isnan(obj) or np.isinf(obj)):
        return None
    return obj

# 退款原因中文映射
REASON_MAP = {
    'No longer needed': '不再需要',
    'Defective item': '产品有缺陷',
    'Item doesn\'t match description': '商品与描述不符',
    'Damaged item or packaging': '商品或包装损坏',
    'Product wouldn\'t arrive on time': '商品不能按时到达',
    'Missing parts': '缺少配件',
    'Missing package': '包裹丢失',
    'Wrong item was sent': '发错商品',
    'Item arrived too late': '商品到达太晚',
    'Package lost': '包裹丢失',
    'Package delivery failed': '配送失败'
}


def load_csv_data(file_path):
    """加载CSV数据"""
    df = pd.read_csv(file_path)
    df['Time Requested'] = pd.to_datetime(
        df['Time Requested'].astype(str).str.strip(),
        format='%d/%m/%Y %H:%M:%S',
        errors='coerce'
    )
    df['Refund Time'] = pd.to_datetime(
        df['Refund Time'].astype(str).str.strip(),
        format='%d/%m/%Y %H:%M:%S',
        errors='coerce'
    )
    df['Order Amount_num'] = pd.to_numeric(
        df['Order Amount'].astype(str).str.replace('$', '').str.replace(',', ''),
        errors='coerce'
    )
    df['Return unit price_num'] = pd.to_numeric(
        df['Return unit price'].astype(str).str.replace('$', '').str.replace(',', ''),
        errors='coerce'
    )
    df['日期'] = df['Time Requested'].apply(
        lambda x: f"{x.year}.{x.month}.{x.day}" if pd.notna(x) else ''
    )
    df['日期排序'] = df['Time Requested'].dt.floor('D')
    df['退款原因（中文）'] = df['Return Reason'].map(REASON_MAP).fillna(df['Return Reason'])
    return df


def analyze_returns(df):
    """执行退货分析"""
    results = {}

    # 总览统计
    total_orders = len(df)
    total_amount = df['Return unit price_num'].sum()
    total_quantity = df['Return Quantity'].sum()
    avg_return_amount = total_amount / total_orders if total_orders > 0 else 0

    results['overview'] = {
        '总退货订单数': int(total_orders),
        '总退货金额': round(float(total_amount), 2),
        '总退货数量': int(total_quantity),
        '平均退货金额': round(float(avg_return_amount), 2),
        '仅退款订单数': int(len(df[df['Return Type'] == 'Refund only'])),
        '退货退款订单数': int(len(df[df['Return Type'] == 'Return and refund']))
    }

    # 按日期统计
    daily_stats = df.groupby(['日期', '日期排序']).agg({
        'Order ID': 'count',
        'Return unit price_num': 'sum',
        'Return Quantity': 'sum'
    }).reset_index()
    daily_stats.columns = ['日期', '日期排序', '退货订单数', '退货金额', '退货数量']
    daily_stats['平均退货金额'] = (daily_stats['退货金额'] / daily_stats['退货订单数']).round(2)
    daily_stats = daily_stats.sort_values('日期排序')
    results['daily'] = daily_stats[['日期', '退货订单数', '退货金额', '退货数量', '平均退货金额']].to_dict('records')

    # 按退货类型统计
    type_stats = df.groupby('Return Type').agg({
        'Order ID': 'count',
        'Return unit price_num': 'sum',
        'Return Quantity': 'sum'
    }).reset_index()
    type_stats.columns = ['退货类型', '订单数', '退货金额', '退货数量']
    type_stats['占比'] = (type_stats['订单数'] / type_stats['订单数'].sum() * 100).round(2)
    type_stats['退货类型'] = type_stats['退货类型'].map({
        'Return and refund': '退货退款',
        'Refund only': '仅退款'
    })
    results['by_type'] = type_stats.to_dict('records')

    # 按退货原因统计
    reason_stats = df.groupby('Return Reason').agg({
        'Order ID': 'count',
        'Return unit price_num': 'sum'
    }).reset_index()
    reason_stats.columns = ['退货原因', '订单数', '退货金额']
    reason_stats['占比'] = (reason_stats['订单数'] / reason_stats['订单数'].sum() * 100).round(2)
    reason_stats['原因中文'] = reason_stats['退货原因'].map(REASON_MAP).fillna(reason_stats['退货原因'])
    reason_stats = reason_stats.sort_values('订单数', ascending=False)
    results['by_reason'] = reason_stats.to_dict('records')

    # 退货订单明细
    detail_cols = ['日期', 'Order ID', 'Order Amount_num', 'Seller SKU', 'Return Type',
                   'Return Reason', '退款原因（中文）', 'Return unit price_num',
                   'Return Quantity', 'Buyer Note']
    detail_df = df[detail_cols].copy()
    detail_df.columns = ['日期', '订单号', '要求退款金额', 'Seller SKU', '退货类型',
                         '退款原因', '退款原因（中文）', '实际退款金额',
                         '退款数量', '退款理由']
    detail_df['订单号'] = detail_df['订单号'].astype(str)
    results['details'] = detail_df.to_dict('records')

    # 月度统计（用于图表）
    df['月份'] = df['Time Requested'].apply(
        lambda x: f"{x.year}.{x.month}" if pd.notna(x) else ''
    )
    monthly_stats = df.groupby('月份').agg({
        'Order ID': 'count',
        'Return unit price_num': 'sum',
        'Return Quantity': 'sum'
    }).reset_index()
    monthly_stats.columns = ['月份', '退货订单数', '退货金额', '退货数量']
    monthly_stats = monthly_stats.sort_values('月份')
    results['monthly'] = monthly_stats.to_dict('records')

    # SKU统计
    sku_stats = df.groupby('Seller SKU').agg({
        'Order ID': 'count',
        'Return unit price_num': 'sum',
        'Return Quantity': 'sum'
    }).reset_index()
    sku_stats.columns = ['SKU', '退货订单数', '退货金额', '退货数量']
    sku_stats = sku_stats.sort_values('退货金额', ascending=False).head(20)
    results['top_skus'] = sku_stats.to_dict('records')

    return results


def load_all_orders(file_path):
    """加载全部订单数据计算退货率"""
    from openpyxl import load_workbook
    order_wb = load_workbook(file_path, data_only=True)
    order_ws = order_wb.active
    order_headers = [cell.value for cell in order_ws[1]]
    order_data = []
    for row in order_ws.iter_rows(min_row=3, values_only=True):
        order_data.append(row)
    df_orders = pd.DataFrame(order_data, columns=order_headers)

    def parse_delivered_time(dt_str):
        if pd.isna(dt_str) or not dt_str:
            return None
        try:
            return pd.to_datetime(str(dt_str).strip(), format='%m/%d/%Y %I:%M:%S %p', errors='coerce')
        except:
            return None

    df_orders['Delivered_dt'] = df_orders['Delivered Time'].apply(parse_delivered_time)
    df_orders = df_orders[df_orders['Delivered_dt'].notna()].copy()
    df_orders['is_return'] = df_orders['Cancelation/Return Type'].isin(['Return/Refund'])

    # 按月统计
    df_orders['month_label'] = df_orders['Delivered_dt'].apply(
        lambda x: x.strftime('%Y.%m') if pd.notna(x) else ''
    )
    monthly_stats = df_orders.groupby('month_label').agg(
        总订单数=('Order ID', 'count'),
        退货订单数=('is_return', 'sum')
    ).reset_index()
    monthly_stats.columns = ['月份', '总订单数', '退货订单数']
    monthly_stats['退货率'] = (monthly_stats['退货订单数'] / monthly_stats['总订单数'] * 100).round(2)
    monthly_stats = monthly_stats[monthly_stats['月份'] != '']

    # 按周统计
    df_orders['week_label'] = df_orders['Delivered_dt'].dt.to_period('W').apply(
        lambda r: f"{r.start_time.strftime('%Y.%m.%d')}_{r.end_time.strftime('%m.%d')}"
    )
    weekly_stats = df_orders.groupby('week_label').agg(
        总订单数=('Order ID', 'count'),
        退货订单数=('is_return', 'sum')
    ).reset_index()
    weekly_stats.columns = ['时段', '总订单数', '退货订单数']
    weekly_stats['退货率'] = (weekly_stats['退货订单数'] / weekly_stats['总订单数'] * 100).round(2)

    return {
        'weekly': weekly_stats.to_dict('records'),
        'monthly_return_rate': monthly_stats.to_dict('records')
    }


def main(version, data_dir, all_orders_file=None):
    """主函数"""
    data_path = os.path.join(data_dir, version, '退货_退款订单.csv')
    output_path = os.path.join(data_dir, version, 'analysis_result.json')

    if not os.path.exists(data_path):
        print(f"数据文件不存在: {data_path}")
        sys.exit(1)

    # 加载并分析
    df = load_csv_data(data_path)
    results = analyze_returns(df)

    # 如果有全部订单数据，添加退货率
    if all_orders_file and os.path.exists(all_orders_file):
        try:
            return_rate_data = load_all_orders(all_orders_file)
            results['return_rate'] = return_rate_data
        except Exception as e:
            print(f"警告: 退货率计算失败 ({all_orders_file}): {e}")
            results['return_rate'] = None

    results['version'] = version
    results['generated_at'] = datetime.now().isoformat()

    # 保存结果
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(clean_for_json(results), f, ensure_ascii=False, indent=2)

    print(f"分析完成: {output_path}")
    return output_path


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("用法: python return_analyzer.py <version> <data_dir> [all_orders_file]")
        sys.exit(1)

    version = sys.argv[1]
    data_dir = sys.argv[2]
    all_orders_file = sys.argv[3] if len(sys.argv) > 3 else None

    main(version, data_dir, all_orders_file)
